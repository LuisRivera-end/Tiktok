from pathlib import Path
from docx import Document


SOURCE = Path(r"C:\Users\lelie\Downloads\Reporte de Estadía Completo (Capítulo I al V, Orange y TikTok).docx")
OUTPUT = Path(r"E:\Tiktok\Reporte_Estadia_CRISP_DM_laboratorio.docx")

# Edición sobre el archivo original: se conservan secciones, estilos, tablas y
# encabezados. El texto reemplazado conserva el formato del primer tramo.
P = {
    8: "ANÁLISIS DE UN PROTOTIPO DE RECOMENDACIÓN DE VIDEO CORTO MEDIANTE CRISP-DM, ORANGE DATA MINING Y PRUEBAS DE LABORATORIO",
    33: "A. Contexto de ByteDance / TikTok y Veta",
    36: "iii. Objetivos del proyecto",
    27: "Este reporte documenta Veta, un prototipo educativo de recomendación de video corto inspirado en problemas de personalización a gran escala. Su implementación local permite explorar recuperación de candidatos, filtros, puntuación multiobjetivo heurística y diversificación. La referencia a TikTok y Monolith ofrece contexto teórico; no implica acceso a los sistemas, datos ni infraestructura de ByteDance.",
    28: "El estudio sigue las seis fases de CRISP-DM y usa un CSV exportado del laboratorio del prototipo. Orange Data Mining se empleó para construir y evaluar un clasificador exploratorio de abandono temprano. Se distinguen las mediciones reproducibles de los objetivos pendientes: el clasificador no está integrado en el feed, y el CSV sintético no demuestra cambios en retención real, equidad de exposición ni latencia de producción.",
    44: "H. Metodología CRISP-DM",
    48: "2.2 Representación Vectorial y Similitud Coseno",
    49: "2.3 Puntuación Multiobjetivo y MMoE como Referente Teórico",
    50: "2.4 Diversificación por Reglas y DPP como Referente Teórico",
    53: "3.2 Comprensión y Preparación de Datos en Orange",
    54: "3.3 Implementación del Prototipo y Modelado Exploratorio",
    60: "4.1 Perfil y Calidad de los Datos de Laboratorio",
    61: "4.2 Evaluación del Clasificador de Abandono Temprano",
    72: "A. Contexto de ByteDance / TikTok y Veta",
    73: "ByteDance y TikTok se estudian como contexto de sistemas de recomendación de video corto. Veta es un desarrollo académico independiente; las pruebas de este informe corresponden exclusivamente al prototipo y a eventos sintéticos generados en su laboratorio.",
    75: "Como contexto, TikTok declara la misión de inspirar creatividad y brindar alegría. La misión operativa de Veta es ofrecer un entorno reproducible para estudiar personalización, señales de consumo y evaluación de modelos.",
    77: "La visión del proyecto académico es disponer de un laboratorio local que permita medir el comportamiento del recomendador y contrastar hipótesis con datos trazables antes de proponer una aplicación real.",
    78: "iii. Objetivos del proyecto",
    79: "Registrar interacciones con una definición clara de evento y unidad de análisis.",
    80: "Examinar abandono temprano, finalización, diversidad y exposición a contenido emergente en datos del laboratorio.",
    81: "Evaluar modelos con particiones independientes y documentar sus límites antes de cualquier integración operativa.",
    83: "La filosofía de trabajo es separar hipótesis, implementación y evidencia. Las cifras del CSV sintético se interpretan como resultados del laboratorio, no como métricas de TikTok ni de usuarios reales.",
    85: "En un feed de video corto, optimizar una sola señal puede favorecer recomendaciones repetitivas o con consumo superficial. Para estudiar este problema, Veta genera interacciones sintéticas y expone un recomendador modular. El problema de minería evaluado aquí es predecir el abandono temprano en eventos terminales, con especial cuidado de evitar variables que revelen la respuesta después de la reproducción.",
    87: "Analizar el prototipo Veta mediante CRISP-DM y evaluar, con datos de laboratorio y Orange Data Mining, si las características disponibles antes de la reproducción permiten distinguir eventos de abandono temprano de eventos de finalización.",
    89: "Describir el pipeline implementado de recuperación, filtrado, puntuación y diversificación, e identificar sus diferencias frente a arquitecturas teóricas como Monolith, MMoE, HNSW y DPP.",
    90: "Definir el objetivo early_skip, la unidad de análisis y las variables excluidas para prevenir fuga de información.",
    91: "Perfilar el CSV de laboratorio y preparar una muestra de eventos terminales con partición por sesión.",
    92: "Comparar modelos de clasificación mediante validación agrupada y prueba final en sesiones reservadas, informando AUC, precisión, sensibilidad, F1 y matriz de confusión.",
    93: "Documentar limitaciones, riesgos de generalización e integración futura del modelo con el prototipo.",
    95: "Veta implementa un pipeline local configurable. Recupera candidatos por similitud coseno, afinidad temática/audio, exploración de publicaciones emergentes o popularidad para usuarios nuevos; aplica filtros; calcula cinco señales heurísticas ponderadas; y diversifica evitando repeticiones contiguas de autor, audio o categoría. Orange analiza el CSV exportado y produce un clasificador separado. HNSW, MMoE y DPP se mantienen como referentes para trabajo futuro, no como componentes ya ejecutados.",
    97: "El valor del proyecto consiste en disponer de un prototipo y una ruta de evaluación verificable. El análisis de datos permite detectar definiciones ambiguas, riesgos de fuga de información y diferencias entre un resultado de clasificación y una mejora del sistema de recomendación. Esta distinción es necesaria para tomar decisiones de ingeniería responsables.",
    100: "Implementación local de Veta con API FastAPI, PostgreSQL, MongoDB y cliente React para simulación, registro de eventos y consulta del feed.",
    101: "Pipeline heurístico configurable con recuperación, filtros, puntuación multiobjetivo y reglas de diversidad.",
    102: "Laboratorio de eventos sintéticos, exportación CSV y análisis reproducible del abandono temprano.",
    103: "Evaluación fuera de muestra del clasificador con sesiones reservadas; comparación con regresión logística como referencia.",
    104: "Identificación de riesgos de calidad, generalización, privacidad y monitoreo antes de integrar el modelo.",
    106: "No se estudian datos ni sistemas internos de ByteDance. Los registros son sintéticos y los resultados no representan usuarios de TikTok.",
    107: "La partición por sesión reduce dependencia entre eventos, pero entrenamiento y prueba comparten el universo de usuarios y videos. No hay validación con usuarios, publicaciones o periodos nuevos.",
    108: "No se ha medido impacto causal del clasificador sobre retención, diversidad, equidad, ingresos, ROI ni latencia de producción; tampoco está desplegado en el ranking del feed.",
    109: "H. Metodología CRISP-DM",
    110: "Se aplicaron las seis fases iterativas de CRISP-DM con un entregable verificable en cada una:",
    111: "Comprensión del negocio: objetivo académico de estudiar abandono temprano y evaluar si existe señal predictiva útil; criterio de éxito: desempeño fuera de muestra y documentación de límites.",
    112: "Comprensión de datos: inspección de esquema, tipos de evento, usuarios, videos, sesiones, faltantes y filas idénticas del CSV exportado por el laboratorio.",
    113: "Preparación de datos: selección de eventos terminales complete y skip; etiqueta early_skip; exclusión de event_type, watch_ms y completion_ratio por fuga de información, y de identificadores directos.",
    114: "Modelado: Random Forest y regresión logística con variables previas a la reproducción. El modelo se entrenó como artefacto de Orange; el puntuador del feed permanece heurístico.",
    115: "Evaluación: validación agrupada por sesión en entrenamiento, selección de umbral y una prueba final en sesiones reservadas; análisis de AUC, precisión, sensibilidad, F1 y errores.",
    116: "Despliegue: se conserva el flujo .ows, las tablas y el modelo serializado para reproducir el experimento. Su incorporación al feed requiere pruebas de integración, calibración y monitoreo.",
    118: "El Capítulo I presenta el problema, alcance y metodología CRISP-DM. El Capítulo II separa los conceptos de recomendación del comportamiento realmente implementado. El Capítulo III describe el software, los datos y el experimento de modelado con Orange.",
    119: "El Capítulo IV comunica los datos observados y la evaluación fuera de muestra. El Capítulo V establece conclusiones limitadas a la evidencia, riesgos y pasos necesarios para futuras pruebas operativas.",
    124: "CRISP-DM organiza la minería de datos en comprensión del negocio, comprensión de datos, preparación, modelado, evaluación y despliegue. En este trabajo la fuente es una exportación de eventos sintéticos del laboratorio de Veta. Cada fila representa un evento, no necesariamente una reproducción independiente; por ello la unidad de modelado se restringe a eventos terminales.",
    125: "Orange Data Mining permite explorar distribuciones, definir atributos y comparar clasificadores de forma visual. Aquí se usó para el flujo de clasificación de abandono temprano; no se atribuyen al experimento análisis PCA, K-Means o correlaciones que no quedaron documentados con datos verificables.",
    126: "2.2 Representación Vectorial y Similitud Coseno",
    127: "Veta representa afinidades del usuario y atributos del video con vectores y calcula similitud coseno para ordenar candidatos. La implementación actual recorre y ordena el catálogo elegible en memoria; no utiliza un índice de búsqueda aproximada HNSW.",
    130: "Monolith es una referencia de investigación sobre recomendación en tiempo real con tablas de embeddings dinámicos. Ese diseño no forma parte del código de Veta ni de las pruebas de este informe. En el prototipo, la similitud coseno es una señal de recuperación y puntuación.",
    131: "2.3 Puntuación Multiobjetivo y MMoE como Referente Teórico",
    133: "MMoE es una arquitectura neuronal para aprendizaje multitarea. Veta no entrena ni ejecuta una red MMoE: calcula cinco salidas mediante fórmulas heurísticas basadas en afinidad, tema, audio y tasa previa de skip. Esta distinción impide interpretar sus salidas como probabilidades calibradas o como pérdidas optimizadas durante entrenamiento.",
    134: "Tabla I. Cinco señales heurísticas del puntuador implementado en Veta",
    135: "2.4 Diversificación por Reglas y DPP como Referente Teórico",
    136: "Los DPP son un método de selección diversa estudiado en la literatura. Veta utiliza una regla codificada más simple: tras ordenar por puntuación, evita que videos contiguos repitan autor, audio o categoría cuando existen alternativas. No se calculó un DPP ni un índice cuantitativo de diversidad en este experimento.",
    141: "El prototipo Veta filtra clips inactivos, bloqueados, vistos y duplicados antes de asignar los cupos de recuperación. Después reaplica filtros, puntúa y diversifica; el reciclaje omite hashes de contenido ya elegidos. Sus parámetros predeterminados son hasta 80 candidatos de origen, 40 tras la ordenación y 20 resultados finales: límites configurados, no mediciones de rendimiento.",
    144: "Catálogo local de videos elegibles (tamaño variable)",
    146: "Etapa 1: Recuperación por similitud, tema/audio, exploración o popularidad → hasta 80 candidatos",
    148: "Etapa 2: Filtros de actividad, autor, vistos, bloqueo, edad y duplicados → tamaño variable",
    150: "Etapa 3: Puntuación heurística multiobjetivo → hasta 40 candidatos ordenados",
    152: "Etapa 4: Regla de diversidad y ensamblado → hasta 20 videos en el feed",
    153: "Fig. 1. Esquema del pipeline local. La elegibilidad se comprueba antes de recuperar y los filtros se reaplican sobre los candidatos.",
    154: "3.2 Comprensión y Preparación de Datos en Orange",
    155: "El archivo veta_interactions.csv contiene 12,271 filas de eventos sintéticos, 17 columnas, 80 usuarios, 210 videos y 524 sesiones. Los eventos complete (3,809) y skip (1,888) forman 5,697 registros terminales. En esta muestra todos los skip cumplen early_skip=1, definido en la exportación como skip con watch_ms menor de 2,000; esto no autoriza a redefinir la etiqueta como cualquier skip en otros conjuntos.",
    156: "Flujo reproducible de datos y modelado en Orange Data Mining:",
    158: "CSV → selección de eventos terminales y revisión de tipos/faltantes",
    160: "Exclusión de fuga de información → partición por session_id → entrenamiento y validación agrupada",
    162: "Random Forest y regresión logística → prueba reservada → matriz de confusión y curva ROC",
    163: "Fig. 2. Flujo de clasificación del abandono temprano; las tablas preparadas y el archivo .ows permiten repetirlo.",
    164: "La tasa de early_skip entre eventos terminales es 33.14% (1,888/5,697). El CSV no contiene un identificador único de evento ni marca temporal: 5,040 filas son idénticas en sus 17 campos, pero no es posible decidir solo con este archivo si representan duplicados indebidos o interacciones coincidentes. campaign_id falta en 12,027 filas y pass_id en 6,271; se excluyeron del modelo. También se excluyeron watch_ms, completion_ratio y event_type porque revelan el desenlace.",
    165: "3.3 Implementación del Prototipo y Modelado Exploratorio",
    167: "La recuperación se realiza sobre el catálogo elegible en memoria. La configuración predeterminada reserva aproximadamente 60% de un máximo de 80 candidatos a similitud coseno, 30% a afinidad temática y 10% a exploración; los usuarios nuevos usan popularidad como alternativa. Los cupos reales dependen de disponibilidad y deduplicación.",
    168: "Canal vectorial: ordena por similitud coseno con el perfil del usuario; no se implementó HNSW ni un tiempo máximo de 8 ms.",
    169: "Canal temático: combina afinidad por categoría, audio y etiquetas existentes en el perfil.",
    170: "Canal exploratorio: prioriza publicaciones con menos de 100 reproducciones en el catálogo local. Esto crea una oportunidad de exposición, sin demostrar equidad de distribución.",
    172: "Los filtros comprueban que el video esté activo y sea elegible para el usuario; descartan contenido propio, visto, bloqueado, restringido por edad o duplicado según los campos disponibles. No hay modelos de análisis visual/textual ni detección perceptual de spam en esta implementación.",
    174: "El puntuador de Veta calcula cinco señales y1 a y5 mediante funciones deterministas; no existe una red MMoE entrenada en este módulo. Las señales representan finalización, tiempo visto, repetición, interacción social y riesgo de abandono como estimaciones heurísticas.",
    175: "Puntuación = 0.35·y1 + 0.30·y2 + 0.15·y3 + 0.12·y4 − 0.08·y5 + ajustes de contexto",
    176: "La configuración aplica una penalización adicional de 0.45 cuando y5 supera 0.70 y pequeños ajustes por coincidencia regional o seguimiento. Estos coeficientes son parámetros del prototipo, no pesos aprendidos ni evidencia de mayor retención.",
    178: "Tras puntuar, Veta conserva hasta 40 candidatos y diversifica hasta 20 resultados. La regla voraz evita repeticiones contiguas de autor, audio o categoría y procura incluir un clip elegible de la región del usuario. No implementa DPP.",
    182: "4.1 Perfil y Calidad de los Datos de Laboratorio",
    183: "La evidencia procede de una exportación puntual del generador de eventos sintéticos del laboratorio, no de una simulación verificada de 30 días con 10,000 usuarios. Ese generador no llama a rank_organic_feed: los cambios recientes de filtrado previo y reciclaje del feed no recalculan este CSV ni las métricas de Orange. La tabla resume conteos observados. No existe un experimento A/B ni medición comparable de retención, ROI o latencia P99.",
    184: "Tabla II. Perfil verificable de la exportación y alcance de las métricas",
    185: "Las categorías de evento son heterogéneas: una reproducción puede generar varias filas (impresión, play, heartbeat, interacción y término). Por ello, 3,809 complete y 1,888 skip se interpretan como eventos terminales; el porcentaje de finalización dentro de ese subconjunto es 66.86%. No debe extrapolarse a una tasa de éxito de usuarios reales ni a una mejora del ranking.",
    186: "4.2 Evaluación del Clasificador de Abandono Temprano",
    187: "Se reservaron 94 sesiones (1,171 eventos terminales) para prueba final y se entrenó con 373 sesiones (4,526 eventos). La selección de umbral 0.24 se hizo con validación agrupada por sesión dentro del entrenamiento. Random Forest obtuvo AUC-ROC 0.850, precisión 0.700, sensibilidad 0.731 y F1 0.715 para early_skip=1 en la prueba reservada. La regresión logística de referencia obtuvo AUC-ROC 0.702 en esa misma prueba.",
    188: "Tabla III. Matriz de confusión de Random Forest en sesiones reservadas (umbral 0.24)",
    189: "La exactitud global fue 78.57% (920/1,171). De los 431 abandonos tempranos reales, 315 se detectaron y 116 no; de las 740 finalizaciones, 605 se clasificaron correctamente y 135 se marcaron como abandono. El AUC mide discriminación en esta muestra; no demuestra calibración ni desempeño en producción.",
    190: "Límites de generalización y decisión de integración:",
    192: "La partición por sesión evita mezclar una misma sesión entre entrenamiento y prueba, pero usuarios, videos, temas y audios pueden reaparecer en ambos grupos.",
    194: "Antes de integrar el modelo: validar con usuarios y videos nuevos, periodo posterior, calibración por región y categoría, deriva de datos, latencia y efectos del feed en una prueba controlada.",
    195: "",
    196: "El archivo CSV no contiene una asignación experimental de usuarios a versiones del recomendador. Por tanto, no se calculan mejoras causales de diversidad, equidad, permanencia o satisfacción. El clasificador aporta una línea de base para investigación, no una validación del sistema de recomendación completo.",
    201: "El proyecto dispone de un prototipo funcional de recomendación heurística y de un flujo reproducible de minería de datos alineado con CRISP-DM. El laboratorio produjo 12,271 eventos sintéticos; 5,697 eventos terminales permitieron evaluar un clasificador de abandono temprano con AUC-ROC 0.850 en sesiones reservadas.",
    202: "El resultado solo respalda la discriminación en esta muestra. No se entrenó MMoE ni se implementaron HNSW o DPP; tampoco se demostró una mejora del feed. El filtrado previo evita gastar cupos en clips inelegibles y el reciclaje impide hashes repetidos, pero su efecto cuantitativo sobre el feed requiere una prueba separada. La ausencia de fecha/ID de evento y el solapamiento de usuarios y videos entre particiones limitan la generalización.",
    204: "Para continuar el trabajo se recomiendan acciones verificables:",
    205: "Añadir event_id, fecha/hora, versión de política de recomendación y run_id a la exportación; validar la semántica de cada evento y conservar un diccionario de datos.",
    206: "Repetir la evaluación con usuarios, videos y periodos realmente nuevos; informar precisión, sensibilidad, F1, AUC y calibración por segmentos, junto con intervalos de incertidumbre.",
    207: "Medir latencia, diversidad, exposición y retención en un experimento controlado antes de afirmar impacto o calcular retorno económico; definir monitoreo de deriva, privacidad y criterios de reversión.",
    209: "El trabajo futuro prioriza la conexión cuidadosa entre análisis y producto:",
    210: "Integrar opcionalmente el clasificador como una señal del puntuador y verificar que no introduzca sesgos por región, categoría, usuario nuevo o creador emergente.",
    211: "Comparar el pipeline heurístico con líneas de base de recomendación bajo el mismo catálogo y la misma asignación experimental; estudiar por separado el módulo publicitario ya presente en Veta.",
    212: "Evaluar métodos más complejos como MMoE, índices aproximados y DPP solo si los datos, las pruebas y el costo operativo justifican su implementación.",
    230: "Repositorio local del proyecto Veta. (2026). Código de API, recomendador, simulador, exportación CSV y documentación técnica. Evidencia inspeccionada en E:\\Tiktok.",
    231: "Orange Data Mining. (2026). Documentación y aplicación de análisis visual. Disponible en: https://orangedatamining.com/ [consulta: septiembre de 2026].",
    232: "CRISP-DM Consortium. (2000). CRISP-DM 1.0: Step-by-step data mining guide. Marco metodológico de seis fases.",
    235: "Apéndice A. Lógica del Puntuador Heurístico Implementado",
    236: "El siguiente resumen corresponde al código local del prototipo y no a una inferencia neuronal de ByteDance:",
    237: "Paso 1: Obtener atributos de usuario, video y contexto; determinar elegibilidad y fuentes de candidatos.",
    238: "Paso 2: Calcular y1 a y5 con funciones heurísticas de afinidad vectorial, tema, audio y tasa previa de skip.",
    239: "Paso 3: Aplicar una penalización de 0.45 si y5 supera el umbral configurado de 0.70.",
    240: "Paso 4: Sumar las señales con pesos 0.35, 0.30, 0.15, 0.12 y −0.08, más ajustes de contexto.",
    241: "Paso 5: Ordenar, seleccionar hasta 40 candidatos y diversificar hasta 20 resultados.",
    245: "Entorno: prototipo local con API FastAPI, PostgreSQL, MongoDB y cliente React. No se verificó un clúster AMD EPYC ni un despliegue distribuido de producción.",
    246: "Datos de esta evaluación: CSV de 12,271 filas y 17 columnas; 80 usuarios, 210 videos, 524 sesiones y 5,697 eventos terminales. No se documentan 10,000 perfiles, 50,000 videos ni 500,000 eventos.",
    247: "Artefactos reproducibles: veta_modelado.ows, veta_terminal_train.tab, veta_terminal_test.tab, veta_early_skip_model.pkcls y veta_early_skip_report.json. La partición final se hizo por sesión; no equivale a una validación en usuarios o videos nuevos.",
}

TABLES = [
    [
        ["Identificador", "Señal heurística", "Cálculo implementado", "Interpretación"],
        ["Y1", "Finalización", "Sigmoide de afinidad, tema, audio y skip previo", "Estimación no calibrada"],
        ["Y2", "Tiempo visto", "Combinación acotada de afinidad, tema y skip", "Razón estimada"],
        ["Y3", "Repetición", "Sigmoide de afinidad y audio", "Estimación no calibrada"],
        ["Y4", "Interacción social", "Sigmoide de afinidad y tema", "Estimación no calibrada"],
        ["Y5", "Abandono temprano", "Sigmoide de afinidad, skip y tema", "Penalización si supera 0.70"],
    ],
    [
        ["Métrica evaluada", "Valor observado", "Unidad / base", "Procedencia", "Alcance"],
        ["Eventos exportados", "12,271", "Filas / 17 columnas", "CSV de laboratorio", "Eventos sintéticos"],
        ["Usuarios y videos", "80 / 210", "IDs distintos", "CSV de laboratorio", "Sin muestra externa"],
        ["Sesiones", "524", "IDs distintos", "CSV de laboratorio", "Sin tiempo en CSV"],
        ["Eventos terminales", "5,697", "3,809 complete + 1,888 skip", "CSV de laboratorio", "Base del modelo"],
        ["Abandono temprano", "33.14%", "1,888 / 5,697 terminales", "CSV de laboratorio", "Solo esta muestra"],
        ["Filas idénticas", "5,040", "17 campos iguales", "CSV de laboratorio", "Sin event_id para resolver"],
    ],
    [
        ["Clase real", "Predicción: finalización (0)", "Predicción: abandono (1)", "Total real"],
        ["Finalización (0)", "605 correctas", "135 falsas alarmas", "740"],
        ["Abandono temprano (1)", "116 no detectados", "315 correctos", "431"],
    ],
]


def replace_paragraph(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main():
    doc = Document(SOURCE)
    for index, value in P.items():
        replace_paragraph(doc.paragraphs[index], value)
    for index in [111, 112, 113, 114, 115, 116, 168, 169, 170, 237, 238, 239, 240, 241, 245, 246, 247]:
        for run in doc.paragraphs[index].runs:
            run.bold = False
    for index in [134, 138, 139, 140, 182, 184, 188]:
        doc.paragraphs[index].paragraph_format.keep_with_next = True
    doc.paragraphs[138].paragraph_format.page_break_before = True
    for table, data in zip(doc.tables, TABLES):
        assert len(table.rows) == len(data)
        for row, values in zip(table.rows, data):
            assert len(row.cells) == len(values)
            for cell, value in zip(row.cells, values):
                replace_paragraph(cell.paragraphs[0], value)
                for extra in cell.paragraphs[1:]:
                    replace_paragraph(extra, "")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
